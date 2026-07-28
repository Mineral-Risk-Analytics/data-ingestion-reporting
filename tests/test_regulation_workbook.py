"""Regulation workbook loader tests (2026-07-23).

Covers the source-of-truth sync semantics: create with provenance, update
with field diff, archive-on-absence, SUGGESTED_* isolation, validation
rejects, weights JSONB rebuild, scope replacement, and dry-run rollback.
"""

from __future__ import annotations

from datetime import date

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker

import app.models  # noqa: F401 — register all ORM models
from app.db.base import Base
from app.models.regulatory import (
    Regulation,
    RegulationGeographyScope,
    RegulationMaterialScope,
)
from app.models.supply import Material
from app.services.ingestion.regulation_workbook import load_regulation_workbook

REG_HDRS = [
    "regulation_key", "title", "issuing_body", "geography", "policy_theme",
    "status", "publication_date", "effective_date", "is_obligation",
    "obligation_points", "verified", "summary", "applies_all_materials",
    "source_url", "source_note", "last_verified_date",
]


def _patch_sqlite_jsonb() -> None:
    from sqlalchemy.dialects.sqlite.base import SQLiteTypeCompiler  # type: ignore[import]
    if not hasattr(SQLiteTypeCompiler, "visit_JSONB"):
        SQLiteTypeCompiler.visit_JSONB = SQLiteTypeCompiler.visit_JSON  # type: ignore[attr-defined]


@pytest.fixture()
def session() -> Session:
    _patch_sqlite_jsonb()
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    s = sessionmaker(bind=engine)()
    s.add(Material(canonical_name="Cobalt"))
    s.add(Material(canonical_name="Gallium"))
    s.commit()
    try:
        yield s
    finally:
        s.close()
        engine.dispose()


def _workbook(regs, weights=(), mats=(), geos=()):
    wb = Workbook()
    ws = wb.active
    ws.title = "Regulations"
    ws.append(REG_HDRS)
    for r in regs:
        ws.append(r)
    ws = wb.create_sheet("GeographyWeights")
    ws.append(["regulation_key", "country_code", "weight"])
    for r in weights:
        ws.append(r)
    ws = wb.create_sheet("MaterialScopes")
    ws.append(["regulation_key", "material", "scope_type", "notes"])
    for r in mats:
        ws.append(r)
    ws = wb.create_sheet("GeographyScopes")
    ws.append(["regulation_key", "country_code", "scope_type"])
    for r in geos:
        ws.append(r)
    return wb


def _reg_row(key, *, pts=10, verified=False, status="effective",
             src="https://example.com/rule", ob=True, all_mats=False):
    return [key, f"{key} title", "Issuer", "US", "theme", status,
            "2025-01-01", "2025-06-01", ob, pts, verified,
            "summary text", all_mats, src, "note", "2026-07-23"]


def _save(tmp_path, wb, name="wb.xlsx"):
    p = tmp_path / name
    wb.save(p)
    return p


class TestCreate:
    def test_creates_with_provenance_weights_scopes(self, session, tmp_path):
        wb = _workbook(
            [_reg_row("TEST_REG")],
            weights=[["TEST_REG", "CN", 1.0], ["TEST_REG", "DEFAULT", 0.1]],
            mats=[["TEST_REG", "Cobalt", "restricted", "why"]],
            geos=[["TEST_REG", "CN", "jurisdiction"]],
        )
        report = load_regulation_workbook(session, _save(tmp_path, wb))
        assert report.created == ["TEST_REG"]
        assert report.rejected == []
        reg = session.scalars(select(Regulation)).one()
        assert reg.obligation_points == 10
        assert reg.is_obligation is True
        assert reg.geography_compliance_weights == {"CN": 1.0, "DEFAULT": 0.1}
        assert reg.metadata_json["source_url"] == "https://example.com/rule"
        assert reg.metadata_json["last_verified"] == "2026-07-23"
        assert reg.publication_date == date(2025, 1, 1)
        scopes = session.scalars(select(RegulationMaterialScope)).all()
        assert [(s.scope_type, s.notes) for s in scopes] == [("restricted", "why")]
        geos = session.scalars(select(RegulationGeographyScope)).all()
        assert [(g.country_code, g.scope_type) for g in geos] == [("CN", "jurisdiction")]

    def test_new_reg_without_source_url_rejected(self, session, tmp_path):
        wb = _workbook([_reg_row("NO_SRC", src=None)])
        report = load_regulation_workbook(session, _save(tmp_path, wb))
        assert report.created == []
        assert report.rejected[0]["reason"] == "new regulation requires source_url"


class TestUpdateAndArchive:
    def _seed(self, session, **kw):
        reg = Regulation(
            regulation_key="SEEDED", title="old title", is_obligation=True,
            obligation_points=5, status="effective",
            metadata_json={"seed_version": 3}, **kw,
        )
        session.add(reg)
        session.commit()
        return reg

    def test_update_diffs_fields_and_preserves_metadata(self, session, tmp_path):
        self._seed(session)
        wb = _workbook([_reg_row("SEEDED", pts=12)])
        report = load_regulation_workbook(session, _save(tmp_path, wb))
        assert report.updated == ["SEEDED"]
        reg = session.scalars(select(Regulation)).one()
        assert reg.obligation_points == 12
        assert reg.title == "SEEDED title"
        assert reg.metadata_json["seed_version"] == 3      # preserved
        assert reg.metadata_json["source_url"] == "https://example.com/rule"

    def test_absent_curated_reg_is_archived(self, session, tmp_path):
        self._seed(session)
        wb = _workbook([_reg_row("OTHER_REG")])
        report = load_regulation_workbook(session, _save(tmp_path, wb))
        assert report.archived == ["SEEDED"]
        seeded = session.scalars(
            select(Regulation).where(Regulation.regulation_key == "SEEDED")
        ).one()
        assert seeded.status == "archived"

    def test_suggested_rows_never_touched(self, session, tmp_path):
        session.add(Regulation(regulation_key="SUGGESTED_FR_1", status="proposed"))
        session.commit()
        wb = _workbook([
            _reg_row("REAL_REG"),
            _reg_row("SUGGESTED_FR_2"),   # someone pasted a discovery key
        ])
        report = load_regulation_workbook(session, _save(tmp_path, wb))
        assert report.created == ["REAL_REG"]
        assert any("SUGGESTED_*" in r["reason"] for r in report.rejected)
        suggested = session.scalars(
            select(Regulation).where(Regulation.regulation_key == "SUGGESTED_FR_1")
        ).one()
        assert suggested.status == "proposed"              # NOT archived

    def test_unchanged_rerun_is_noop(self, session, tmp_path):
        wb = _workbook(
            [_reg_row("TEST_REG")],
            weights=[["TEST_REG", "CN", 1.0]],
            mats=[["TEST_REG", "Cobalt", "covered", None]],
        )
        p = _save(tmp_path, wb)
        load_regulation_workbook(session, p)
        report2 = load_regulation_workbook(session, p)
        assert report2.unchanged == ["TEST_REG"]
        assert report2.updated == []


class TestValidation:
    def test_bad_rows_rejected_rest_load(self, session, tmp_path):
        wb = _workbook(
            [
                _reg_row("GOOD"),
                _reg_row("BAD_STATUS", status="bogus"),
                _reg_row("BAD_PTS", pts=99),
                _reg_row("OB_NO_PTS", pts=None),
            ],
            weights=[
                ["GOOD", "CN", 1.5],          # out of range
                ["GOOD", "CHN", 0.5],         # bad ISO2
                ["GOOD", "CN", 0.9],
            ],
            mats=[
                ["GOOD", "Unobtainium", "covered", None],   # unknown material
                ["GOOD", "Cobalt", "bogus_scope", None],    # bad scope type
                ["GOOD", "Gallium", "restricted", None],
            ],
        )
        report = load_regulation_workbook(session, _save(tmp_path, wb))
        assert report.created == ["GOOD"]
        reasons = [r["reason"] for r in report.rejected]
        assert "invalid status 'bogus'" in reasons
        assert "obligation_points 99 outside [0, 40]" in reasons
        assert "is_obligation=TRUE requires obligation_points >= 1" in reasons
        assert any("outside [0, 1]" in r for r in reasons)
        assert any("invalid country_code" in r for r in reasons)
        assert any("unknown material" in r for r in reasons)
        assert any("invalid scope_type" in r for r in reasons)
        reg = session.scalars(select(Regulation)).one()
        assert reg.geography_compliance_weights == {"CN": 0.9}
        scopes = session.scalars(select(RegulationMaterialScope)).all()
        assert len(scopes) == 1

    def test_duplicate_key_rejected(self, session, tmp_path):
        wb = _workbook([_reg_row("DUP"), _reg_row("DUP")])
        report = load_regulation_workbook(session, _save(tmp_path, wb))
        assert report.created == ["DUP"]
        assert any("duplicate" in r["reason"] for r in report.rejected)


class TestDryRun:
    def test_dry_run_rolls_back(self, session, tmp_path):
        wb = _workbook([_reg_row("TEST_REG")])
        report = load_regulation_workbook(session, _save(tmp_path, wb), dry_run=True)
        assert report.created == ["TEST_REG"]
        assert report.dry_run is True
        assert session.scalars(select(Regulation)).all() == []


class TestRegionAliasKeys:
    def test_eu_region_weight_row_accepted(self, session, tmp_path):
        wb = _workbook(
            [_reg_row("REGIONAL")],
            weights=[["REGIONAL", "EU", 0.0], ["REGIONAL", "DEFAULT", 0.3]],
        )
        report = load_regulation_workbook(session, _save(tmp_path, wb))
        assert report.rejected == []
        reg = session.scalars(select(Regulation)).one()
        assert reg.geography_compliance_weights == {"EU": 0.0, "DEFAULT": 0.3}


class TestAppliesAllMaterialsFlag:
    def test_flag_round_trips(self, session, tmp_path):
        wb = _workbook([_reg_row("ALL_GOODS", all_mats=True), _reg_row("SCOPED")])
        report = load_regulation_workbook(session, _save(tmp_path, wb))
        assert report.rejected == []
        regs = {r.regulation_key: r for r in session.scalars(select(Regulation)).all()}
        assert regs["ALL_GOODS"].applies_all_materials is True
        assert regs["SCOPED"].applies_all_materials is False


class TestEditorialSheets:
    def _editorial_wb(self):
        wb = _workbook([_reg_row("ED_REG")])
        ws = wb.create_sheet("Editorial")
        ws.append(["regulation_key", "standfirst", "what_it_requires",
                   "who_must_comply", "materials_and_origins", "key_dates",
                   "why_it_matters"])
        ws.append(["ED_REG", "The standfirst.", "Requires X.", "Importers.",
                   "Cobalt from CD.", "Since 2024.", "It matters."])
        ws2 = wb.create_sheet("FurtherReading")
        ws2.append(["regulation_key", "title", "publisher", "url"])
        ws2.append(["ED_REG", "Official text", "EUR-Lex", "https://example.com/text"])
        ws2.append(["ED_REG", "Bad link", "X", "not-a-url"])
        return wb

    def test_editorial_synced_to_metadata(self, session, tmp_path):
        report = load_regulation_workbook(session, _save(tmp_path, self._editorial_wb()))
        reg = session.scalars(select(Regulation)).one()
        ed = reg.metadata_json["editorial"]
        assert ed["standfirst"] == "The standfirst."
        assert ed["sections"]["why_it_matters"] == "It matters."
        assert ed["further_reading"] == [
            {"title": "Official text", "publisher": "EUR-Lex", "url": "https://example.com/text"}
        ]
        assert any("url must be absolute" in r["reason"] for r in report.rejected)

    def test_workbook_without_editorial_sheets_still_loads(self, session, tmp_path):
        wb = _workbook([_reg_row("PLAIN")])
        report = load_regulation_workbook(session, _save(tmp_path, wb))
        assert report.created == ["PLAIN"]
        assert report.rejected == []


class TestEnforcementWeightsSheet:
    def test_enforcement_sheet_round_trips(self, session, tmp_path):
        wb = _workbook([_reg_row("ENF_REG", all_mats=True)])
        ws = wb.create_sheet("EnforcementWeights")
        ws.append(["regulation_key", "material", "weight"])
        ws.append(["ENF_REG", "Cobalt", 0.6])
        ws.append(["ENF_REG", "DEFAULT", 0.3])
        ws.append(["ENF_REG", "Unobtainium", 0.5])   # unknown -> rejected
        ws.append(["ENF_REG", "Gallium", 1.5])        # out of range -> rejected
        report = load_regulation_workbook(session, _save(tmp_path, wb))
        reg = session.scalars(select(Regulation)).one()
        assert reg.material_enforcement_weights == {"Cobalt": 0.6, "DEFAULT": 0.3}
        reasons = [r["reason"] for r in report.rejected]
        assert any("unknown material" in r for r in reasons)
        assert any("outside [0, 1]" in r for r in reasons)

    def test_absent_sheet_leaves_null(self, session, tmp_path):
        wb = _workbook([_reg_row("NOENF_REG")])
        load_regulation_workbook(session, _save(tmp_path, wb))
        reg = session.scalars(select(Regulation)).one()
        assert reg.material_enforcement_weights is None
