"""Partner-curated regulation workbook loader (2026-07-23).

Reads the multi-sheet regulation curation workbook and syncs the DB to it.
The workbook is the SOURCE OF TRUTH for curated regulations:

  - ``Regulations``      → upserts ``regulations`` rows by ``regulation_key``
  - ``GeographyWeights`` → rebuilds ``geography_compliance_weights`` JSONB
  - ``MaterialScopes``   → replaces ``regulation_material_scope`` rows
  - ``GeographyScopes``  → replaces ``regulation_geography_scope`` rows

Sync semantics (decided 2026-07-23, Nicole):
  * Full sync for curated regulations: a curated DB regulation (key not
    LIKE ``SUGGESTED_%``) missing from the workbook is ARCHIVED
    (``status='archived'``) — never hard-deleted, and its scopes/weights
    are left in place so un-archiving restores it.
  * Discovery suggestions (``SUGGESTED_*``) are ignored entirely: never
    written by this loader, never archived by it, and rejected with a
    warning if someone adds one to the workbook.
  * Provenance is required: new regulations must carry ``source_url``.
    ``source_url`` / ``source_note`` / ``last_verified_date`` are stored
    in ``metadata_json`` (keys ``source_url``, ``source_note``,
    ``last_verified``) without clobbering other metadata.

Validation mirrors seed_facilities_partner.py: a bad row is rejected and
reported; the rest of the workbook still loads. ``dry_run=True`` builds
the full diff report and rolls back.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

import structlog
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.regulatory import (
    Regulation,
    RegulationGeographyScope,
    RegulationMaterialScope,
)
from app.models.supply import Material

log = structlog.get_logger(__name__)

VALID_STATUSES = frozenset({"effective", "enacted", "proposed", "stayed", "archived"})
VALID_MATERIAL_SCOPE_TYPES = frozenset(
    {"banned", "restricted", "strategic_raw_material", "covered", "disclosure_required"}
)
VALID_GEO_SCOPE_TYPES = frozenset({"jurisdiction", "targeted_country"})

_REG_SHEET = "Regulations"
_WEIGHT_SHEET = "GeographyWeights"
_MAT_SHEET = "MaterialScopes"
_GEO_SHEET = "GeographyScopes"
# Optional content-site sheets (2026-07-23): editorial standfirst + sectioned
# body, and a structured further-reading list. Synced to
# metadata_json["editorial"] = {standfirst, sections{...}, further_reading[...]}.
_EDITORIAL_SHEET = "Editorial"
_FR_SHEET = "FurtherReading"
# Optional (2026-07-27): per-material enforcement weights for the obligation
# uplift (regulation_key | material | weight) → material_enforcement_weights
# JSONB keyed by canonical name (+ DEFAULT). Absent sheet/rows = NULL = 1.0.
_ENF_SHEET = "EnforcementWeights"
_EDITORIAL_SECTIONS = (
    "what_it_requires", "who_must_comply", "materials_and_origins",
    "key_dates", "why_it_matters",
)

# Regulation columns synced verbatim from the sheet.
_SYNCED_FIELDS = (
    "title", "issuing_body", "geography", "policy_theme", "status",
    "publication_date", "effective_date", "is_obligation",
    "obligation_points", "verified", "summary", "applies_all_materials",
)


@dataclass
class WorkbookReport:
    created: list[str] = field(default_factory=list)
    updated: list[str] = field(default_factory=list)
    unchanged: list[str] = field(default_factory=list)
    archived: list[str] = field(default_factory=list)
    rejected: list[dict[str, Any]] = field(default_factory=list)
    dry_run: bool = False

    def to_dict(self) -> dict[str, Any]:
        return {
            "created": self.created,
            "updated": self.updated,
            "unchanged": self.unchanged,
            "archived": self.archived,
            "rejected": self.rejected,
            "dry_run": self.dry_run,
            "counts": {
                "created": len(self.created),
                "updated": len(self.updated),
                "unchanged": len(self.unchanged),
                "archived": len(self.archived),
                "rejected": len(self.rejected),
            },
        }


def _reject(report: WorkbookReport, sheet: str, row: int, key: Any, reason: str) -> None:
    report.rejected.append(
        {"sheet": sheet, "row": row, "regulation_key": key, "reason": reason}
    )


def _cell_str(v: Any) -> Optional[str]:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _cell_bool(v: Any) -> Optional[bool]:
    if isinstance(v, bool):
        return v
    s = _cell_str(v)
    if s is None:
        return None
    if s.upper() in ("TRUE", "YES", "1"):
        return True
    if s.upper() in ("FALSE", "NO", "0"):
        return False
    return None


def _cell_date(v: Any) -> Optional[date]:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _cell_str(v)
    if s is None:
        return None
    try:
        return date.fromisoformat(s[:10])
    except ValueError:
        return None


def _read_sheet(wb, name: str) -> tuple[list[str], list[tuple[int, tuple]]]:
    """Return (headers, [(excel_row_number, values), ...]) skipping blank rows."""
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [(_cell_str(h) or "") for h in rows[0]]
    out = []
    for i, vals in enumerate(rows[1:], start=2):
        if all(v is None or _cell_str(v) is None for v in vals):
            continue
        out.append((i, vals))
    return headers, out


def _col(headers: list[str], name: str) -> int:
    try:
        return headers.index(name)
    except ValueError as exc:
        raise ValueError(f"workbook missing required column '{name}'") from exc


def load_regulation_workbook(
    session: Session,
    xlsx_path: str | Path,
    *,
    dry_run: bool = False,
) -> WorkbookReport:
    """Sync the regulations tables to the workbook. See module docstring."""
    report = WorkbookReport(dry_run=dry_run)
    wb = load_workbook(str(xlsx_path), data_only=True)
    for required in (_REG_SHEET, _WEIGHT_SHEET, _MAT_SHEET, _GEO_SHEET):
        if required not in wb.sheetnames:
            raise ValueError(f"workbook missing required sheet '{required}'")

    materials = {
        m.canonical_name: m.id
        for m in session.scalars(select(Material)).all()
    }

    # ── Pass 1: weights / scopes keyed by regulation_key ────────────────
    hdrs, rows = _read_sheet(wb, _WEIGHT_SHEET)
    w_key, w_cc, w_wt = _col(hdrs, "regulation_key"), _col(hdrs, "country_code"), _col(hdrs, "weight")
    weights_by_key: dict[str, dict[str, float]] = {}
    for rn, vals in rows:
        key, cc = _cell_str(vals[w_key]), _cell_str(vals[w_cc])
        if not key or not cc:
            _reject(report, _WEIGHT_SHEET, rn, key, "missing regulation_key or country_code")
            continue
        cc = cc.upper()
        # Valid keys: ISO2 country, a region alias (constants.REGION_MEMBERS,
        # e.g. "EU"), or DEFAULT.
        from app.constants import REGION_MEMBERS
        if cc != "DEFAULT" and cc not in REGION_MEMBERS and (
            len(cc) != 2 or not cc.isalpha()
        ):
            _reject(report, _WEIGHT_SHEET, rn, key, f"invalid country_code '{cc}'")
            continue
        try:
            wt = float(vals[w_wt])
        except (TypeError, ValueError):
            _reject(report, _WEIGHT_SHEET, rn, key, "weight is not a number")
            continue
        if not 0.0 <= wt <= 1.0:
            _reject(report, _WEIGHT_SHEET, rn, key, f"weight {wt} outside [0, 1]")
            continue
        weights_by_key.setdefault(key, {})[cc] = wt

    hdrs, rows = _read_sheet(wb, _MAT_SHEET)
    m_key, m_mat, m_st = _col(hdrs, "regulation_key"), _col(hdrs, "material"), _col(hdrs, "scope_type")
    m_notes = _col(hdrs, "notes")
    mat_scopes_by_key: dict[str, list[tuple[int, str, Optional[str]]]] = {}
    for rn, vals in rows:
        key, mat, st = _cell_str(vals[m_key]), _cell_str(vals[m_mat]), _cell_str(vals[m_st])
        if not key or not mat:
            _reject(report, _MAT_SHEET, rn, key, "missing regulation_key or material")
            continue
        if st not in VALID_MATERIAL_SCOPE_TYPES:
            _reject(report, _MAT_SHEET, rn, key, f"invalid scope_type '{st}'")
            continue
        if mat not in materials:
            _reject(report, _MAT_SHEET, rn, key, f"unknown material '{mat}'")
            continue
        mat_scopes_by_key.setdefault(key, []).append(
            (materials[mat], st, _cell_str(vals[m_notes]))
        )

    hdrs, rows = _read_sheet(wb, _GEO_SHEET)
    g_key, g_cc, g_st = _col(hdrs, "regulation_key"), _col(hdrs, "country_code"), _col(hdrs, "scope_type")
    geo_scopes_by_key: dict[str, list[tuple[str, str]]] = {}
    for rn, vals in rows:
        key, cc, st = _cell_str(vals[g_key]), _cell_str(vals[g_cc]), _cell_str(vals[g_st])
        if not key or not cc:
            _reject(report, _GEO_SHEET, rn, key, "missing regulation_key or country_code")
            continue
        cc = cc.upper()
        if len(cc) != 2 or not cc.isalpha():
            _reject(report, _GEO_SHEET, rn, key, f"invalid country_code '{cc}'")
            continue
        if st not in VALID_GEO_SCOPE_TYPES:
            _reject(report, _GEO_SHEET, rn, key, f"invalid scope_type '{st}'")
            continue
        geo_scopes_by_key.setdefault(key, []).append((cc, st))

    # ── Pass 1a½: optional enforcement-weights sheet ────────────────────
    enf_by_key: dict[str, dict[str, float]] = {}
    if _ENF_SHEET in wb.sheetnames:
        hdrs, rows = _read_sheet(wb, _ENF_SHEET)
        e_key = _col(hdrs, "regulation_key")
        e_mat, e_wt = _col(hdrs, "material"), _col(hdrs, "weight")
        for rn, vals in rows:
            key, mat = _cell_str(vals[e_key]), _cell_str(vals[e_mat])
            if not key or not mat:
                _reject(report, _ENF_SHEET, rn, key, "missing regulation_key or material")
                continue
            if mat != "DEFAULT" and mat not in materials:
                _reject(report, _ENF_SHEET, rn, key, f"unknown material '{mat}'")
                continue
            try:
                wt = float(vals[e_wt])
            except (TypeError, ValueError):
                _reject(report, _ENF_SHEET, rn, key, "weight is not a number")
                continue
            if not 0.0 <= wt <= 1.0:
                _reject(report, _ENF_SHEET, rn, key, f"weight {wt} outside [0, 1]")
                continue
            enf_by_key.setdefault(key, {})[mat] = wt

    # ── Pass 1b: optional editorial sheets ──────────────────────────────
    editorial_by_key: dict[str, dict[str, Any]] = {}
    if _EDITORIAL_SHEET in wb.sheetnames:
        hdrs, rows = _read_sheet(wb, _EDITORIAL_SHEET)
        e_key = _col(hdrs, "regulation_key")
        e_stand = _col(hdrs, "standfirst")
        e_cols = {sec: _col(hdrs, sec) for sec in _EDITORIAL_SECTIONS}
        for rn, vals in rows:
            key = _cell_str(vals[e_key])
            if not key:
                _reject(report, _EDITORIAL_SHEET, rn, None, "missing regulation_key")
                continue
            standfirst = _cell_str(vals[e_stand])
            sections = {
                sec: txt for sec in _EDITORIAL_SECTIONS
                if (txt := _cell_str(vals[e_cols[sec]])) is not None
            }
            if not standfirst and not sections:
                continue
            entry: dict[str, Any] = {}
            if standfirst:
                entry["standfirst"] = standfirst
            if sections:
                entry["sections"] = sections
            editorial_by_key[key] = entry

    further_by_key: dict[str, list[dict[str, str]]] = {}
    if _FR_SHEET in wb.sheetnames:
        hdrs, rows = _read_sheet(wb, _FR_SHEET)
        f_key = _col(hdrs, "regulation_key")
        f_title, f_pub, f_url = _col(hdrs, "title"), _col(hdrs, "publisher"), _col(hdrs, "url")
        for rn, vals in rows:
            key, title, url = _cell_str(vals[f_key]), _cell_str(vals[f_title]), _cell_str(vals[f_url])
            if not key or not title or not url:
                _reject(report, _FR_SHEET, rn, key, "further-reading row needs regulation_key, title, url")
                continue
            if not url.startswith(("http://", "https://")):
                _reject(report, _FR_SHEET, rn, key, f"url must be absolute http(s): '{url}'")
                continue
            further_by_key.setdefault(key, []).append({
                "title": title,
                "publisher": _cell_str(vals[f_pub]) or "",
                "url": url,
            })

    # ── Pass 2: Regulations sheet upsert ────────────────────────────────
    hdrs, rows = _read_sheet(wb, _REG_SHEET)
    cols = {name: _col(hdrs, name) for name in (
        "regulation_key", *_SYNCED_FIELDS, "source_url", "source_note", "last_verified_date",
    )}
    existing = {
        r.regulation_key: r
        for r in session.scalars(select(Regulation)).all()
    }
    workbook_keys: set[str] = set()

    for rn, vals in rows:
        key = _cell_str(vals[cols["regulation_key"]])
        if not key:
            _reject(report, _REG_SHEET, rn, None, "missing regulation_key")
            continue
        if key.startswith("SUGGESTED_"):
            _reject(report, _REG_SHEET, rn, key,
                    "SUGGESTED_* keys are discovery rows — curate under a real key instead")
            continue
        if key in workbook_keys:
            _reject(report, _REG_SHEET, rn, key, "duplicate regulation_key in workbook")
            continue

        status = _cell_str(vals[cols["status"]])
        if status is not None and status not in VALID_STATUSES:
            _reject(report, _REG_SHEET, rn, key, f"invalid status '{status}'")
            continue
        is_ob = _cell_bool(vals[cols["is_obligation"]])
        verified = _cell_bool(vals[cols["verified"]])
        pts_raw = vals[cols["obligation_points"]]
        pts: Optional[int] = None
        if _cell_str(pts_raw) is not None:
            try:
                pts = int(float(pts_raw))
            except (TypeError, ValueError):
                _reject(report, _REG_SHEET, rn, key, "obligation_points is not an integer")
                continue
            if pts < 0 or pts > 40:
                _reject(report, _REG_SHEET, rn, key, f"obligation_points {pts} outside [0, 40]")
                continue
        if is_ob and not pts:
            _reject(report, _REG_SHEET, rn, key,
                    "is_obligation=TRUE requires obligation_points >= 1")
            continue

        source_url = _cell_str(vals[cols["source_url"]])
        reg = existing.get(key)
        if reg is None and not source_url:
            _reject(report, _REG_SHEET, rn, key, "new regulation requires source_url")
            continue

        new_vals: dict[str, Any] = {
            "title": _cell_str(vals[cols["title"]]),
            "issuing_body": _cell_str(vals[cols["issuing_body"]]),
            "geography": _cell_str(vals[cols["geography"]]),
            "policy_theme": _cell_str(vals[cols["policy_theme"]]),
            "status": status,
            "publication_date": _cell_date(vals[cols["publication_date"]]),
            "effective_date": _cell_date(vals[cols["effective_date"]]),
            "is_obligation": bool(is_ob),
            "obligation_points": pts,
            "verified": bool(verified),
            "summary": _cell_str(vals[cols["summary"]]),
            "applies_all_materials": bool(
                _cell_bool(vals[cols["applies_all_materials"]])
            ),
        }
        workbook_keys.add(key)

        if reg is None:
            reg = Regulation(regulation_key=key, **new_vals)
            session.add(reg)
            session.flush()
            report.created.append(key)
            changed = True
        else:
            changed = False
            for f_ in _SYNCED_FIELDS:
                if getattr(reg, f_) != new_vals[f_]:
                    setattr(reg, f_, new_vals[f_])
                    changed = True

        # Provenance merged into metadata_json (other keys preserved).
        meta = dict(reg.metadata_json or {})
        prov = {
            "source_url": source_url,
            "source_note": _cell_str(vals[cols["source_note"]]),
            "last_verified": (
                d.isoformat() if (d := _cell_date(vals[cols["last_verified_date"]])) else None
            ),
        }
        prov = {k: v for k, v in prov.items() if v is not None}
        if prov and {k: meta.get(k) for k in prov} != prov:
            meta.update(prov)
            reg.metadata_json = meta
            changed = True

        # Editorial block (content site): replace wholesale from the sheets.
        editorial = dict(editorial_by_key.get(key) or {})
        if further_by_key.get(key):
            editorial["further_reading"] = further_by_key[key]
        if editorial:
            if meta.get("editorial") != editorial:
                meta = dict(reg.metadata_json or {})
                meta["editorial"] = editorial
                reg.metadata_json = meta
                changed = True

        # Weights JSONB rebuilt from the long-format sheet.
        new_w = weights_by_key.get(key) or None
        if reg.geography_compliance_weights != new_w:
            reg.geography_compliance_weights = new_w
            changed = True

        # Enforcement weights (065): rebuilt the same way; None = 1.0 for all.
        new_enf = enf_by_key.get(key) or None
        if reg.material_enforcement_weights != new_enf:
            reg.material_enforcement_weights = new_enf
            changed = True

        # Scopes: replace wholesale (sheet is authoritative).
        want_mat = sorted(mat_scopes_by_key.get(key, []))
        have_mat = sorted(
            (s.material_id, s.scope_type, s.notes)
            for s in session.scalars(
                select(RegulationMaterialScope).where(
                    RegulationMaterialScope.regulation_id == reg.id
                )
            ).all()
        )
        if want_mat != have_mat:
            for s in session.scalars(
                select(RegulationMaterialScope).where(
                    RegulationMaterialScope.regulation_id == reg.id
                )
            ).all():
                session.delete(s)
            for mat_id, st, notes in want_mat:
                session.add(RegulationMaterialScope(
                    regulation_id=reg.id, material_id=mat_id,
                    scope_type=st, notes=notes,
                ))
            changed = True

        want_geo = sorted(geo_scopes_by_key.get(key, []))
        have_geo = sorted(
            (s.country_code, s.scope_type)
            for s in session.scalars(
                select(RegulationGeographyScope).where(
                    RegulationGeographyScope.regulation_id == reg.id
                )
            ).all()
        )
        if want_geo != have_geo:
            for s in session.scalars(
                select(RegulationGeographyScope).where(
                    RegulationGeographyScope.regulation_id == reg.id
                )
            ).all():
                session.delete(s)
            for cc, st in want_geo:
                session.add(RegulationGeographyScope(
                    regulation_id=reg.id, country_code=cc, scope_type=st,
                ))
            changed = True

        if key not in report.created:
            (report.updated if changed else report.unchanged).append(key)

    # ── Pass 3: archive curated DB regs absent from the workbook ────────
    for key, reg in existing.items():
        if key.startswith("SUGGESTED_"):
            continue
        if key not in workbook_keys and reg.status != "archived":
            reg.status = "archived"
            report.archived.append(key)

    session.flush()
    if dry_run:
        session.rollback()
    else:
        session.commit()

    log.info("regulation_workbook.load.done", **report.to_dict()["counts"], dry_run=dry_run)
    return report
