"""Benchmark production-share workbook loader.

Loads human-audited, citation-backed production shares into
``hs_code_production_shares`` for HS nodes that USGS MCS does not cover
(midstream / battery-grade stages).  Companion to the coverage gap
analysis in ``docs/design/concentration_coverage_gaps.md`` (2026-07-15):
after the L1 concentration-purity filter, share coverage is the SOLE
source of the Material Concentration pillar, so these rows are what
makes midstream concentration (CN refined cobalt, ID nickel
intermediates, ...) visible at L1.

Contract (mirrors the manual-risk-events workbook discipline):

* Every row carries a ``source_org`` + ``source_url`` + ``basis`` — the
  loader never invents or derives numbers.
* Rows load ONLY when ``audited == "Y"`` — a human has verified the
  figure against the cited source.  Unaudited rows are counted and
  skipped, so the workbook can be pre-filled for review.
* ``source`` column in the DB becomes ``benchmark_<slug(source_org)>``
  (<= 32 chars), so benchmark rows NEVER collide with ``usgs_mcs`` rows
  on the unique key (mapping, country, year, scope, source) and MCS
  re-runs never touch them.
* Idempotent: insert-if-missing; existing benchmark rows are updated
  only with ``force=True`` (same semantics as the MCS ingester).
* market_scope is always ``global`` — benchmark shares answer "who
  produces globally", never "where do US imports come from".

Newer-year visibility warning
-----------------------------
``hs_node_scorer`` computes HHI from shares at the LATEST reference_year
per node.  If a node already has global rows at a year newer than the
benchmark row being loaded, the benchmark row would be invisible to the
scorer.  The loader loads it anyway (history is fine) but reports it
under ``warnings_shadowed_by_newer_year`` so the operator notices.

Workbook format (sheet "Benchmark Shares", header row 1)
--------------------------------------------------------
material | hs_codes | country_code | production_share | reference_year |
basis | source_org | source_url | notes | audited

``hs_codes`` is a ``;``-separated list of hs_code_prefix values; the row
fans out to one DB row per resolved mapping.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

import structlog
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.country import Country
from app.models.supply import HsCodeMaterialMapping, HsCodeProductionShare, Material

log = structlog.get_logger(__name__)

_SHEET_DEFAULT = "Benchmark Shares"
_COLUMNS = [
    "material", "hs_codes", "country_code", "production_share",
    "reference_year", "basis", "source_org", "source_url", "notes",
    "audited",
]
# Per-node share sums above this trigger a validation error (small
# headroom for independent-source rounding).
_NODE_SUM_TOLERANCE = 1.05


def _slug_source(source_org: str) -> str:
    """'Cobalt Institute' -> 'benchmark_cobalt_institute' (<= 32 chars)."""
    slug = re.sub(r"[^a-z0-9]+", "_", source_org.lower()).strip("_")
    return ("benchmark_" + slug)[:32].rstrip("_")


@dataclass
class BenchmarkShareReport:
    rows_seen: int = 0
    rows_skipped_unaudited: int = 0
    rows_errored: int = 0
    db_rows_inserted: int = 0
    db_rows_updated: int = 0
    db_rows_skipped_existing: int = 0
    errors: list[str] = field(default_factory=list)
    warnings_shadowed_by_newer_year: list[str] = field(default_factory=list)
    dry_run: bool = False

    def as_dict(self) -> dict:
        return {
            "rows_seen": self.rows_seen,
            "rows_skipped_unaudited": self.rows_skipped_unaudited,
            "rows_errored": self.rows_errored,
            "db_rows_inserted": self.db_rows_inserted,
            "db_rows_updated": self.db_rows_updated,
            "db_rows_skipped_existing": self.db_rows_skipped_existing,
            "errors": self.errors,
            "warnings_shadowed_by_newer_year": self.warnings_shadowed_by_newer_year,
            "dry_run": self.dry_run,
        }


def _is_template_hint_row(values: dict) -> bool:
    """Content-based template-row skip (lesson from the facility loader:
    never skip by position — partners delete example rows).

    2026-07-15 bugfix: the original heuristic also skipped any row whose
    combined text exceeded 200 chars — which silently swallowed EVERY
    real row, because the citation ``notes`` column is deliberately long
    (rows_seen=0 on Nicole's first dry run).  Hint detection now looks
    ONLY at the identity columns (material / hs_codes / country_code),
    which real rows keep short and template rows fill with guidance
    text.  The notes column is never inspected.
    """
    ident = " ".join(
        str(values.get(k) or "") for k in ("material", "hs_codes", "country_code")
    ).lower()
    return (
        "e.g." in ident
        or "must match" in ident
        or len(ident) > 80
    )


def seed_benchmark_shares(
    session: Session,
    xlsx_path: str,
    *,
    sheet: str = _SHEET_DEFAULT,
    dry_run: bool = False,
    force: bool = False,
) -> BenchmarkShareReport:
    report = BenchmarkShareReport(dry_run=dry_run)

    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(
            f"Sheet {sheet!r} not found in {xlsx_path} "
            f"(sheets: {wb.sheetnames})"
        )
    ws = wb[sheet]

    rows_iter = ws.iter_rows(values_only=True)
    header = [str(h).strip().lower() if h is not None else "" for h in next(rows_iter)]
    missing_cols = [c for c in _COLUMNS if c not in header]
    if missing_cols:
        raise ValueError(f"Workbook missing columns: {missing_cols}")
    idx = {c: header.index(c) for c in _COLUMNS}

    # ── Preload reference data ────────────────────────────────────────────
    materials = {
        m.canonical_name: m.id
        for m in session.scalars(select(Material)).all()
    }
    countries = set(session.scalars(select(Country.iso2)).all())
    mappings = {
        (mp.material_id, mp.hs_code_prefix): mp
        for mp in session.scalars(
            select(HsCodeMaterialMapping).where(
                HsCodeMaterialMapping.market_scope == "global"
            )
        ).all()
    }

    # First pass: parse + validate all rows (all-or-nothing per file, like
    # the manual-events loader — one bad row should be fixed, not half-loaded).
    parsed: list[dict] = []
    per_node_year_sums: dict[tuple[int, int], float] = {}

    for row_num, raw in enumerate(rows_iter, start=2):
        values = {c: raw[idx[c]] if idx[c] < len(raw) else None for c in _COLUMNS}
        if all(v is None or str(v).strip() == "" for v in values.values()):
            continue
        if _is_template_hint_row(values):
            continue
        report.rows_seen += 1

        audited = str(values["audited"] or "").strip().upper()
        if audited != "Y":
            report.rows_skipped_unaudited += 1
            continue

        errs: list[str] = []
        material = str(values["material"] or "").strip()
        material_id = materials.get(material)
        if material_id is None:
            errs.append(f"unknown material {material!r}")

        cc = str(values["country_code"] or "").strip().upper()
        if cc not in countries:
            errs.append(f"unknown country_code {cc!r}")

        try:
            share = float(values["production_share"])
            if not (0.0 < share <= 1.0):
                errs.append(f"production_share {share} outside (0, 1]")
        except (TypeError, ValueError):
            share = None
            errs.append(f"production_share {values['production_share']!r} not numeric")

        try:
            ref_year = int(values["reference_year"])
            if not (2000 <= ref_year <= 2100):
                errs.append(f"reference_year {ref_year} implausible")
        except (TypeError, ValueError):
            ref_year = None
            errs.append(f"reference_year {values['reference_year']!r} not an integer")

        source_org = str(values["source_org"] or "").strip()
        source_url = str(values["source_url"] or "").strip()
        basis = str(values["basis"] or "").strip()
        if not source_org:
            errs.append("source_org is required")
        if not source_url:
            errs.append("source_url is required (citation-backed rows only)")
        if not basis:
            errs.append("basis is required (what denominator does this share use?)")

        resolved: list[HsCodeMaterialMapping] = []
        hs_codes = [
            c.strip() for c in re.split(r"[;,]", str(values["hs_codes"] or ""))
            if c.strip()
        ]
        if not hs_codes:
            errs.append("hs_codes is empty")
        elif material_id is not None:
            for code in hs_codes:
                mp = mappings.get((material_id, code))
                if mp is None:
                    errs.append(
                        f"no global mapping for ({material!r}, {code!r})"
                    )
                else:
                    resolved.append(mp)

        if errs:
            report.rows_errored += 1
            report.errors.append(f"row {row_num}: " + "; ".join(errs))
            continue

        for mp in resolved:
            key = (mp.id, ref_year)
            per_node_year_sums[key] = per_node_year_sums.get(key, 0.0) + share
            parsed.append({
                "mapping": mp,
                "country_code": cc,
                "production_share": share,
                "reference_year": ref_year,
                "source": _slug_source(source_org),
                "notes": " | ".join(
                    p for p in (
                        f"basis: {basis}",
                        f"source: {source_org}",
                        source_url,
                        str(values["notes"] or "").strip() or None,
                    ) if p
                ),
                "row_num": row_num,
            })

    for (mapping_id, yr), total in sorted(per_node_year_sums.items()):
        if total > _NODE_SUM_TOLERANCE:
            report.rows_errored += 1
            report.errors.append(
                f"node mapping_id={mapping_id} year={yr}: audited shares "
                f"sum to {total:.3f} > {_NODE_SUM_TOLERANCE}"
            )

    if report.errors:
        log.warning(
            "seed_benchmark_shares.validation_failed",
            errors=len(report.errors),
        )
        return report

    # ── Newer-year shadowing check ────────────────────────────────────────
    for p in parsed:
        newest = session.scalar(
            select(HsCodeProductionShare.reference_year)
            .where(
                HsCodeProductionShare.hs_mapping_id == p["mapping"].id,
                HsCodeProductionShare.market_scope == "global",
                HsCodeProductionShare.reference_year > p["reference_year"],
            )
            .order_by(HsCodeProductionShare.reference_year.desc())
            .limit(1)
        )
        if newest is not None:
            report.warnings_shadowed_by_newer_year.append(
                f"row {p['row_num']}: {p['mapping'].hs_code_prefix} "
                f"{p['country_code']} year {p['reference_year']} is shadowed "
                f"by existing year {newest} — scorer uses latest year only"
            )

    # ── Upsert ────────────────────────────────────────────────────────────
    for p in parsed:
        existing = session.scalar(
            select(HsCodeProductionShare).where(
                HsCodeProductionShare.hs_mapping_id == p["mapping"].id,
                HsCodeProductionShare.country_code == p["country_code"],
                HsCodeProductionShare.reference_year == p["reference_year"],
                HsCodeProductionShare.market_scope == "global",
                HsCodeProductionShare.source == p["source"],
            )
        )
        if existing is not None:
            if force:
                existing.production_share = p["production_share"]
                existing.notes = p["notes"]
                report.db_rows_updated += 1
            else:
                report.db_rows_skipped_existing += 1
            continue
        session.add(HsCodeProductionShare(
            hs_mapping_id=p["mapping"].id,
            country_code=p["country_code"],
            production_share=p["production_share"],
            reference_year=p["reference_year"],
            market_scope="global",
            source=p["source"],
            notes=p["notes"],
        ))
        report.db_rows_inserted += 1

    if dry_run:
        session.rollback()
    else:
        session.commit()
    log.info("seed_benchmark_shares.done", **{
        k: v for k, v in report.as_dict().items()
        if not isinstance(v, list)
    })
    return report
